"""Auditable, atomic Mandalay XLSX import followed by CPCB recalculation."""
import argparse
import hashlib
import math
import sys
from datetime import datetime, timezone
from pathlib import Path
from zoneinfo import ZoneInfo

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))
from openpyxl import load_workbook
from app import CITY, FIELDS, POLLUTANTS, TZ, create_app, upsert
from cpcb import mass_convert

MANDALAY_MAPPING = {
    'timestamp':'timestamp', 'original_aqi':'aqi', 'pm2_5':'pm25', 'pm10':'pm10',
    'carbon_monoxide':'co_ug_m3', 'nitrogen_dioxide':'no2_ug_m3',
    'sulphur_dioxide':'so2_ug_m3', 'ozone':'o3_ug_m3',
}


def import_file(path, mapping, *, units, timestamp_timezone, sheet=None,
                confirm=False, source_aqi_standard='unknown'):
    if not confirm:
        raise ValueError('Confirm that the file contains Mandalay data with --confirm-mandalay')
    if 'timestamp' not in mapping or set(mapping)-{'timestamp','original_aqi',*FIELDS}:
        raise ValueError('Supply a timestamp mapping and recognized field names')
    if source_aqi_standard not in ('unknown','us-epa','cpcb'):
        raise ValueError('Unknown source AQI standard')
    if 'us_aqi' in mapping and source_aqi_standard != 'us-epa':
        raise ValueError('An explicit us_aqi mapping requires verified source AQI standard us-epa')
    for field in POLLUTANTS:
        if field in mapping:
            if field not in units:
                raise ValueError(f'Missing verified input unit for {field}')
            mass_convert(1, units[field])  # Validate even when the column is empty.
    zone = ZoneInfo(timestamp_timezone)
    digest = hashlib.sha256(Path(path).read_bytes()).hexdigest()
    source_id = 'xlsx:' + digest
    book = load_workbook(path, read_only=True, data_only=True)
    try:
        ws = book[sheet] if sheet else book.active
        iterator = ws.iter_rows(values_only=True)
        headers = next(iterator)
        if len(set(headers)) != len(headers):
            raise ValueError('Duplicate Excel column headers')
        indices = {field:headers.index(column) for field,column in mapping.items()}
        rows, snapshots, seen = [], [], set()
        for number,cells in enumerate(iterator,2):
            if all(c is None for c in cells):
                continue
            stamp = cells[indices['timestamp']]
            if isinstance(stamp,str):
                stamp = datetime.fromisoformat(stamp)
            if not isinstance(stamp,datetime):
                raise ValueError(f'Row {number}: timestamp must be an Excel datetime or ISO datetime')
            at = (stamp.replace(tzinfo=zone) if stamp.tzinfo is None else stamp).astimezone(timezone.utc)
            local = at.astimezone(TZ)
            if local.minute or local.second or local.microsecond:
                raise ValueError(f'Row {number}: expected a whole-hour Asia/Yangon observation; do not guess timestamp shifts')
            if at in seen:
                raise ValueError(f'Row {number}: duplicate timestamp')
            seen.add(at)
            row = dict(timestamp=at, source='mandalay-xlsx', **{field:None for field in FIELDS})
            for field in POLLUTANTS:
                if field in indices:
                    row[field] = mass_convert(cells[indices[field]], units[field])
            for field in ('us_aqi','original_aqi'):
                if field in indices and cells[indices[field]] is not None:
                    value = float(cells[indices[field]])
                    if not math.isfinite(value) or value < 0:
                        raise ValueError(f'Row {number}: invalid original AQI')
                    if source_aqi_standard == 'us-epa':
                        row['us_aqi'] = value  # Retained raw only; excluded from calculation.
            original = {str(header):(value.isoformat() if isinstance(value,datetime) else value)
                        for header,value in zip(headers,cells)}
            if any(isinstance(v,float) and not math.isfinite(v) for v in original.values()):
                raise ValueError(f'Row {number}: nonfinite original value')
            snapshots.append(dict(timestamp=at, source_id=source_id, original_values=original,
                provenance={'city':CITY, 'workbook_sha256':digest, 'sheet':ws.title, 'row':number,
                            'timestamp_timezone':timestamp_timezone, 'input_units':units,
                            'column_mapping':mapping, 'source_aqi_standard':source_aqi_standard}))
            rows.append(row)
        # Validate every row before starting a single transaction. A DB error
        # rolls back raw measurements, provenance and CPCB results together.
        upsert(rows, snapshots)
        return len(rows)
    finally:
        book.close()


if __name__ == '__main__':
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument('file')
    parser.add_argument('--sheet')
    parser.add_argument('--mandalay-workbook', action='store_true', help='Use the inspected workbook header mapping')
    parser.add_argument('--column', action='append', default=[], help='FIELD=Excel header')
    parser.add_argument('--unit', action='append', default=[], help='FIELD=ug/m3 or mg/m3')
    parser.add_argument('--pm-unit', choices=['ug/m3','mg/m3'], help='Verified PM unit for the inspected workbook preset')
    parser.add_argument('--timestamp-timezone', required=True, help='Verified timezone for naive timestamps')
    parser.add_argument('--confirm-mandalay', action='store_true')
    parser.add_argument('--source-aqi-standard', choices=['unknown','us-epa','cpcb'], default='unknown')
    args = parser.parse_args()
    mapping = dict(MANDALAY_MAPPING) if args.mandalay_workbook else {}
    mapping.update(dict(item.split('=',1) for item in args.column))
    units = {}
    if args.mandalay_workbook:
        if not args.pm_unit:
            parser.error('--mandalay-workbook requires explicit --pm-unit')
        units = {field:'ug/m3' for field in POLLUTANTS}
        units.update(pm2_5=args.pm_unit, pm10=args.pm_unit)
    units.update(dict(item.split('=',1) for item in args.unit))
    with create_app().app_context():
        count = import_file(args.file, mapping, units=units, timestamp_timezone=args.timestamp_timezone,
                            sheet=args.sheet, confirm=args.confirm_mandalay,
                            source_aqi_standard=args.source_aqi_standard)
        print(f'Imported {count} original Mandalay rows; CPCB history recalculated independently')
